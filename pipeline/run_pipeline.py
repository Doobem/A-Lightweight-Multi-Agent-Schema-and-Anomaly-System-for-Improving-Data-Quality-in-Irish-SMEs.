import os
import sys
import pandas as pd
from io import BytesIO

# Ensure project root is in path
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from Integration.graph_sharepoint import (
    download_file_bytes,
    upload_file_to_sharepoint,
    list_files_in_sharepoint_folder,
    move_file_in_sharepoint
)
from Integration.graph_auth import get_access_token
from pipeline.graph_pipeline import build_graph

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

INPUT_FOLDER = "SME_Data/"
OUTPUT_FOLDER = "Documents/AnomalyReports/"
PROCESSED_FOLDER = "Documents/Processed/"
HIGHLIGHTED_FOLDER = "Documents/DetectedError/"

# -------------------------
# HIGHLIGHT FUNCTION
# -------------------------
def highlight_anomalies_in_excel(local_file_path, anomalies, highlighted_path):
    import math

    wb = load_workbook(local_file_path)
    ws = wb[wb.sheetnames[0]]

    yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    orange = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
    red = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

    header_map = {cell.value: idx for idx, cell in enumerate(ws[1], start=1)}

    for anomaly in anomalies:
        row = anomaly["row"]
        col_name = anomaly["column"]
        issue = anomaly["issue_type"]

        # Duplicate row anomalies (Column=None) → highlight entire row
        if col_name is None:
            if row is None:
                continue
            if isinstance(row, float) and math.isnan(row):
                continue
            r = int(row)
            for c in range(1, ws.max_column + 1):
                ws.cell(row=r, column=c).fill = yellow
            continue

        if col_name not in header_map:
            continue

        col_index = header_map[col_name]

        # Invalid row → highlight header
        if row is None or row == "" or (isinstance(row, float) and math.isnan(row)):
            header_cell = ws.cell(row=1, column=col_index)
            header_cell.fill = yellow
            continue

        r = int(row)
        cell = ws.cell(row=r, column=col_index)

        if issue.startswith("Schema/"):
            cell.fill = yellow
        elif issue.startswith("Anomaly/Outlier") or issue.startswith("Anomaly/RareCategory"):
            cell.fill = orange
        elif issue.startswith("Anomaly/Logical"):
            cell.fill = red
        else:
            cell.fill = yellow

    wb.save(highlighted_path)

# -------------------------
# PROCESS SINGLE FILE
# -------------------------

   

def process_file(file_path: str, token: str):
    print(f"\nProcessing: {file_path}\n")

    # 1. Download REAL Excel file bytes
    file_bytes = download_file_bytes(file_path, token)

    # Load DataFrame for anomaly/schema detection
    df = pd.read_excel(BytesIO(file_bytes))

    # Save REAL Excel file locally (basename only)
    local_file_path = os.path.basename(file_path)
    with open(local_file_path, "wb") as f:
        f.write(file_bytes)

    # 2. Run agents
    graph = build_graph()
    result = graph.invoke(
        {"df": df, "generated_anomalies": [], "done": False},
        config={"reset": True}
    )

    anomalies = result["generated_anomalies"]
    schema_results = result.get("schema_results", {})

    expected_columns = ["row", "column", "issue_type", "description", "severity", "category"]

    if not anomalies:
        anomaly_df = pd.DataFrame(columns=expected_columns)
    else:
        anomaly_df = pd.DataFrame(anomalies).reindex(columns=expected_columns)

    dataset_name = os.path.splitext(os.path.basename(file_path))[0]
    output_filename = f"anomaly_report_{dataset_name}.xlsx"

    # 3. Write anomaly report
    with pd.ExcelWriter(output_filename) as writer:

        # ---- FIXED SCHEMA SUMMARY ----
        schema_rows = []
        for col, meta in schema_results.items():
            schema_rows.append({
                "Column": col,
                "Type": meta.get("type"),
                "Missing": meta.get("missing"),
                "Categories": ", ".join(meta.get("categories", []))
                    if isinstance(meta.get("categories"), list)
                    else meta.get("categories")
            })

        pd.DataFrame(schema_rows).to_excel(writer, sheet_name="Schema Summary", index=False)

        # ---- ANOMALIES SHEET ----
        anomaly_df.to_excel(writer, sheet_name="Anomalies", index=False)

    # Convert anomalies to dict list
    anomaly_records = anomaly_df.to_dict(orient="records")

    # 4. Create highlighted dataset
    highlighted_name = dataset_name + "_highlighted.xlsx"
    highlight_anomalies_in_excel(local_file_path, anomaly_records, highlighted_name)

    # Upload highlighted dataset
    upload_file_to_sharepoint(f"{HIGHLIGHTED_FOLDER}{highlighted_name}", highlighted_name, token)

    # 5. Create processed dataset (original renamed)
    processed_name = dataset_name + "_processed.xlsx"
    with open(processed_name, "wb") as f:
        f.write(file_bytes)

    upload_file_to_sharepoint(f"{PROCESSED_FOLDER}{processed_name}", processed_name, token)

    # 6. Move original file in SharePoint
    move_file_in_sharepoint(file_path, f"{PROCESSED_FOLDER}{processed_name}", token)

    # 7. Upload anomaly report
    upload_file_to_sharepoint(f"{OUTPUT_FOLDER}{output_filename}", output_filename, token)

    print(f"Finished: {file_path}")
    print(f"Report uploaded: {OUTPUT_FOLDER}{output_filename}")
    print(f"Processed dataset uploaded: {PROCESSED_FOLDER}{processed_name}")
    print(f"Highlighted dataset uploaded: {HIGHLIGHTED_FOLDER}{highlighted_name}")



# -------------------------
# MAIN
# -------------------------
def main():
    token = get_access_token()

    print("Checking for new datasets in SME_Data...")

    files = list_files_in_sharepoint_folder(INPUT_FOLDER, token)
    new_files = [f for f in files if not f.endswith("_processed.xlsx")]

    if not new_files:
        print("No new files found.")
        return

    print("Checking for new datasets in SME_Data...\n")

    for file_path in new_files:
        process_file(file_path, token)

    print("\nAll new files processed.")

if __name__ == "__main__":
    main()
